import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import clean
import clustering
import report
import texturr


def blobs(n_per=20, centers=((10, 0), (-5, 8.7), (-5, -8.7)), seed=0):   # three distinct directions
    rng = np.random.default_rng(seed)
    X = np.vstack([rng.normal(c, 0.5, (n_per, 2)) for c in centers])
    y = np.repeat(range(len(centers)), n_per)
    return X, y


# --- clustering -------------------------------------------------------------------------------

def test_choose_k_recovers_obvious_structure():
    X, _ = blobs()
    k, scores = clustering.choose_k(X)
    assert k == 3 and scores[3] == max(scores.values())


def test_choose_k_tiny_inputs():
    assert clustering.choose_k(np.zeros((1, 2)))[0] == 1
    assert clustering.choose_k(np.zeros((3, 2)))[0] == 1
    labels, k, _ = clustering.cluster(np.zeros((2, 2)), 'auto')
    assert k == 1 and list(labels) == [0, 0]


def test_cluster_fixed_k_is_capped_by_sample_count():
    X, _ = blobs(n_per=1)
    _, k, _ = clustering.cluster(X, 10)
    assert k == 3


def test_parse_clusters():
    assert clustering.parse_clusters('AUTO') == 'auto' and clustering.parse_clusters('7') == 7
    for bad in ('0', '-2', 'many'):
        with pytest.raises(ValueError):
            clustering.parse_clusters(bad)


# --- reading and selecting -------------------------------------------------------------------

def make_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_positions_survive_blank_leading_row_and_gaps(tmp_path):
    p = str(tmp_path / 'a.xlsx')
    make_xlsx(p, [[None, None], ['Comments', 'Score'], ['one', 1], [None, 2], ['  ', 3], ['two', 4]])
    df = texturr.read_frame(p, None)
    data, pos = texturr.extract_answers(df, 'column', 'A', header_row=2)
    assert data == ['one', 'two']
    out = str(tmp_path / 'out.xlsx')
    report.annotate_workbook(p, None, [{'header': 'Comments', 'positions': pos, 'clusters': [0, 1],
                                        'labels': {0: 'First', 1: 'Second'}}], out, header_row=2)
    ws = load_workbook(out).active
    assert ws['C2'].value == 'Comments Cluster' and ws['D2'].value == 'Comments Theme'
    assert (ws['A3'].value, ws['D3'].value) == ('one', 'First')      # aligned to the right sheet rows
    assert (ws['A6'].value, ws['D6'].value) == ('two', 'Second')
    assert ws['D4'].value is None                                    # skipped blank rows stay untouched


def test_annotate_never_overwrites_data_without_header_row(tmp_path):
    p = str(tmp_path / 'a.xlsx')
    make_xlsx(p, [['first answer'], ['second answer']])
    out = str(tmp_path / 'o.xlsx')
    report.annotate_workbook(p, None, [{'header': 'X', 'positions': [0, 1], 'clusters': [0, 0], 'labels': {0: 'T'}}],
                             out, header_row=0)
    ws = load_workbook(out).active
    assert [ws['A1'].value, ws['A2'].value] == ['first answer', 'second answer']
    assert [ws['B1'].value, ws['B2'].value] == [0, 0]                # cluster ids beside the data
    assert [ws['C1'].value, ws['C2'].value] == ['T', 'T']            # themes; no header cells were written


def test_workbook_theme_cannot_become_a_formula(tmp_path):
    p = str(tmp_path / 'a.xlsx')
    make_xlsx(p, [['Comments'], ['hello']])
    out = str(tmp_path / 'o.xlsx')
    evil = '=HYPERLINK("http://evil","x")'
    report.annotate_workbook(p, None, [{'header': 'Comments', 'positions': [1], 'clusters': [0], 'labels': {0: evil}}], out)
    cell = load_workbook(out).active['C2']
    assert cell.data_type == 's' and cell.value == evil


def test_csv_read_and_annotate_roundtrip(tmp_path):
    p = str(tmp_path / 'a.csv')
    open(p, 'w').write('Comments,Score\nslow app,1\n\nlove it,5\n')
    df = texturr.read_frame(p)
    data, pos = texturr.extract_answers(df, 'column', 'A', 1)
    assert data == ['slow app', 'love it'] and pos == [1, 3]
    out = str(tmp_path / 'o.csv')
    report.annotate_csv(p, [{'header': 'Comments', 'positions': pos, 'clusters': [0, 1],
                             'labels': {0: '=cmd', 1: 'Praise'}}], out, texturr.csv_safe)
    lines = open(out).read().splitlines()
    assert lines[0] == 'Comments,Score,Comments Cluster,Comments Theme'
    assert lines[1].endswith(",0,'=cmd") and lines[3].endswith(',1,Praise')


def test_multi_column_selection_by_name_and_letter():
    df = pd.DataFrame([['Likes', 'Dislikes', 'Score'], ['a', 'b', '1']])
    assert texturr.get_row_or_column(df, 1, column=['dislikes', 'A', 'b'])  == ('column', ['B', 'A'])


def test_args_accept_multiple_columns_and_auto_clusters():
    a = texturr.parse_arguments(['f.xlsx', '--column', 'Likes', 'Dislikes'])
    assert a.column == ['Likes', 'Dislikes'] and a.clusters == 'auto'
    assert texturr.parse_arguments(['f.xlsx', '--clusters', '4']).clusters == 4
    with pytest.raises(SystemExit):
        texturr.parse_arguments(['f.xlsx', '--clusters', 'zero'])


# --- HTML report -----------------------------------------------------------------------------

def test_html_escapes_everything_and_has_no_scripts_or_external_assets():
    rows = [{'Cluster': 0, 'Size': 2, 'Label': '<script>alert(1)</script>', 'Summary': '"quoted" & <b>',
             'Suggested Action': '<img src=x onerror=1>', 'Keyphrases': 'a, <i>b</i>',
             'Representative Responses': '<svg onload=1> | plain', 'Responses': '1, 2'}]
    out = report.render_html('/secret/path/in.xlsx', [('<Col>', rows, 2)], 'by <model>')
    assert '<script' not in out.lower().replace("<script>alert", '') and '&lt;script&gt;' in out
    assert '<img' not in out and '<svg' not in out and '<b>' not in out
    assert 'http://' not in out and 'https://' not in out
    assert 'secret/path' not in out and 'in.xlsx' in out          # only the file name, not the full path
    assert "default-src 'none'" in out


def test_html_sorts_largest_group_first_and_handles_missing_label():
    mk = lambda cid, size, label: {'Cluster': cid, 'Size': size, 'Label': label, 'Summary': '', 'Suggested Action': '',
                                   'Keyphrases': '', 'Representative Responses': '', 'Responses': ''}
    out = report.render_html('f.xlsx', [('C', [mk(0, 1, ''), mk(1, 9, 'Big')], 10)])
    assert out.index('Big') < out.index('Group 0')


# --- end to end with real embeddings (skipped when the ML stack is not installed) --------------------

def test_end_to_end_pipeline(tmp_path):
    pytest.importorskip('sentence_transformers')
    pytest.importorskip('keybert')
    rows = [['Comments']] + [[t] for t in [
        'The app is too slow to load', 'Loading takes forever', 'Pages freeze and crash',
        'Love the new dashboard design', 'Beautiful clean interface', 'Great layout and colors',
        'Support never answered my ticket', 'Customer service was rude', 'Nobody replied to my email']]
    p = str(tmp_path / 'in.xlsx')
    make_xlsx(p, rows)
    args = texturr.parse_arguments([p, '--column', 'Comments', '--llm', 'none', '--clusters', '3'])
    df = texturr.read_frame(p, None)
    data, _ = texturr.extract_answers(df, 'column', 'A', 1)
    out, clusters = texturr.analyze(data, args, None)
    groups = sorted(sorted(i for i, c in enumerate(clusters) if c == k) for k in set(clusters))
    assert groups == [[0, 1, 2], [3, 4, 5], [6, 7, 8]]
    assert all(r['Keyphrases'] for r in out)


# --- non-answers ---------------------------------------------------------------------------

def test_nonanswer_detection():
    for t in ['N/A', ' none ', 'n/a.', '...', '???', 'No comment', "I don\u2019t know", 'idk!', '   ', '-']:
        assert clean.is_nonanswer(t), t
    for t in ['no', 'The app is great', 'None of the buttons work', 'nothing works since the update', 'a']:
        assert not clean.is_nonanswer(t), t


def test_analyze_sets_aside_nonanswers_and_keeps_original_numbering(monkeypatch):
    emb = {'slow': [10, 0], 'slower': [10, 1], 'crash': [10, 0.5], 'nice': [0, 10], 'pretty': [1, 10], 'clean': [0.5, 10], 'N/A': [5, 5], 'none': [5, 5.5]}
    monkeypatch.setattr(texturr, 'compute_embeddings', lambda data, m: np.array([emb[t] for t in data], dtype=float))
    monkeypatch.setattr(texturr, 'extract_keyphrases', lambda d, c, m: {})
    args = texturr.parse_arguments(['f.csv', '--llm', 'none', '--clusters', '2'])
    data = ['slow', 'N/A', 'slower', 'nice', 'none', 'pretty', 'crash', 'clean']
    rows, clusters = texturr.analyze(data, args, None)
    assert list(clusters[[1, 4]]) == [-1, -1] and set(clusters[[0, 2, 6]]) != set(clusters[[3, 5, 7]])
    na = [r for r in rows if r['Cluster'] == -1][0]
    assert na['Label'] == 'Non-answers' and na['Responses'] == '2, 5'
    real = sorted(r['Responses'] for r in rows if r['Cluster'] != -1)
    assert real == ['1, 3, 7', '4, 6, 8']                 # numbers refer to the original rows, not the filtered list
    kept = texturr.analyze(data, texturr.parse_arguments(['f.csv', '--llm', 'none', '--clusters', '2', '--keep-nonanswers']), None)[1]
    assert -1 not in set(kept)


def test_literal_na_text_survives_reading_but_empty_cells_do_not(tmp_path):
    for name, writer in (('a.xlsx', lambda p: make_xlsx(p, [['Comments'], ['N/A'], ['None'], [None], ['real answer'], ['NA']])),
                         ('a.csv', lambda p: open(p, 'w').write('Comments\nN/A\nNone\n\nreal answer\nNA\n'))):
        p = str(tmp_path / name)
        writer(p)
        data, _ = texturr.extract_answers(texturr.read_frame(p, None), 'column', 'A', 1)
        assert data == ['N/A', 'None', 'real answer', 'NA'], name


# --- label evaluation harness -----------------------------------------------------------------------

def test_label_eval_scoring_separates_good_from_bad_labels():
    import sys
    sys.path.insert(0, 'eval')
    import label_eval
    vocab = {'card arrival': [1, 0, 0], 'lost card': [0, 1, 0], 'exchange rate': [0, 0, 1],
             'where is my card': [.9, .1, 0], 'card stolen': [.1, .9, 0], 'fx question': [0, .1, .9]}
    embed = lambda xs: np.array([vocab[x] for x in xs], dtype=float)
    intents = ['card_arrival', 'lost_card', 'exchange_rate']
    top1, sim = label_eval.score(['where is my card', 'card stolen', 'fx question'], intents, embed)
    assert top1 == 1.0 and sim > 0.9
    top1, sim = label_eval.score(['card stolen', 'fx question', 'where is my card'], intents, embed)   # shuffled
    assert top1 == 0.0 and sim < 0.2
    assert label_eval.humanize('card_arrival-time') == 'card arrival time'
