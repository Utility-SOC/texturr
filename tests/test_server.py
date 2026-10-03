import http.client
import json
import threading

import numpy as np
import pytest

import server

TOKEN = 'a' * 32
DIRS = {'slow': (10, 0), 'crash': (9, 2), 'design': (0, 10), 'support': (-8, -8)}


def embed(texts):
    rng = np.random.default_rng(abs(hash(tuple(texts))) % 2**32)
    return np.array([np.array(DIRS[t.split()[0]]) + rng.normal(0, .3, 2) if t.split()[0] in DIRS else [1, 1] for t in texts])


TEXTS = [f'{k} comment {i}' for k in DIRS for i in range(6)] + ['N/A', '=HYPERLINK("http://x")']


@pytest.fixture
def api():
    srv = server.make_server(server.App(TOKEN, embed), '127.0.0.1', 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()

    def call(method, path, body=None, token=TOKEN, raw=False):
        conn = http.client.HTTPConnection('127.0.0.1', srv.server_port, timeout=10)
        headers = {'Authorization': f'Bearer {token}'} if token is not None else {}
        data = None
        if body is not None:
            data = json.dumps(body); headers['Content-Type'] = 'application/json'
        conn.request(method, path, data, headers)
        r = conn.getresponse()
        payload = r.read()
        conn.close()
        return r.status, (payload if raw else (json.loads(payload) if payload else None)), r
    yield call
    srv.shutdown()


def mcp(api, method, params=None, rid=1):
    msg = {'jsonrpc': '2.0', 'method': method, 'params': params or {}}
    if rid is not None:
        msg['id'] = rid
    return api('POST', '/mcp', msg)


def tool(api, name, **args):
    status, out, _ = mcp(api, 'tools/call', {'name': name, 'arguments': args})
    assert status == 200
    res = out['result']
    return res['isError'], (json.loads(res['content'][0]['text']) if not res['isError'] else res['content'][0]['text'])


def make_session(api, clusters=4):
    status, out, _ = api('POST', '/v1/sessions', {'texts': TEXTS, 'clusters': clusters, 'name': 'Test'})
    assert status == 201
    return out


def test_health_is_open_but_everything_else_needs_the_token(api):
    assert api('GET', '/healthz', token=None)[0] == 200
    for token in (None, 'wrong', TOKEN[:-1]):
        assert api('POST', '/v1/sessions', {'texts': ['a']}, token=token)[0] == 401
        assert api('POST', '/mcp', {'jsonrpc': '2.0', 'id': 1, 'method': 'ping'}, token=token)[0] == 401
    status, _, resp = api('GET', '/v1/sessions/x', token=None)
    assert status == 401 and resp.getheader('WWW-Authenticate') == 'Bearer'


def test_create_session_and_overview(api):
    o = make_session(api)
    assert o['groups'] == 4 and o['non_answers'] == 1 and o['responses'] == len(TEXTS) - 0
    status, again, _ = api('GET', f"/v1/sessions/{o['session_id']}")
    assert status == 200 and again['groups'] == 4


def test_create_validation(api):
    for body in ({}, {'texts': []}, {'texts': 'abc'}, {'texts': ['a'], 'clusters': 'many'}):
        status, out, _ = api('POST', '/v1/sessions', body)
        assert status == 400 and 'error' in out


def test_unknown_routes_and_sessions(api):
    assert api('GET', '/nope')[0] == 404
    assert api('GET', '/v1/sessions/doesnotexist')[0] == 400
    assert api('GET', '/mcp')[0] == 405


def test_mcp_handshake_and_tool_list(api):
    status, out, _ = mcp(api, 'initialize', {'protocolVersion': '2025-03-26', 'capabilities': {}, 'clientInfo': {'name': 'n8n'}})
    r = out['result']
    assert status == 200 and r['protocolVersion'] == '2025-03-26' and 'tools' in r['capabilities'] and r['instructions']
    assert mcp(api, 'initialize', {'protocolVersion': '1999-01-01'})[1]['result']['protocolVersion'] == server.SUPPORTED_MCP[0]
    status, body, resp = mcp(api, 'notifications/initialized', rid=None)
    assert status == 202 and body is None
    names = [t['name'] for t in mcp(api, 'tools/list')[1]['result']['tools']]
    assert names == ['get_overview', 'get_cluster_examples', 'merge_clusters', 'split_cluster', 'move_responses', 'set_cluster_label', 'finish_analysis']
    assert mcp(api, 'nonsense')[1]['error']['code'] == -32601


def test_agent_flow_over_mcp_then_report(api):
    sid = make_session(api)['session_id']
    err, o = tool(api, 'get_overview', session_id=sid)
    assert not err and len(o['clusters']) == 4
    c = o['clusters'][0]['id']
    err, ex = tool(api, 'get_cluster_examples', session_id=sid, cluster_id=c, n=2)
    assert not err and len(ex['examples']) == 2
    pair = o['hints']['most_similar_pairs'][0]['clusters']
    err, m = tool(api, 'merge_clusters', session_id=sid, cluster_ids=pair)
    assert not err and m['needs_label']
    err, msg = tool(api, 'finish_analysis', session_id=sid)
    assert err and 'still need labels' in msg                     # cannot finish until everything is labeled
    for g in tool(api, 'get_overview', session_id=sid)[1]['clusters']:
        assert not tool(api, 'set_cluster_label', session_id=sid, cluster_id=g['id'], label=f"=Theme {g['id']}",
                        summary='s', suggested_action='a')[0]
    assert tool(api, 'finish_analysis', session_id=sid)[1]['finished']
    status, rep, _ = api('GET', f'/v1/sessions/{sid}/report?format=json')
    assert status == 200 and [h['op'] for h in rep['history']][:2] == ['created', 'merge'] and rep['history'][-1]['op'] == 'finish'
    status, csv, _ = api('GET', f'/v1/sessions/{sid}/report?format=csv', raw=True)
    assert status == 200 and b"'=Theme" in csv and b',=Theme' not in csv   # label starting with '=' neutralized
    status, html, resp = api('GET', f'/v1/sessions/{sid}/report?format=html', raw=True)
    assert status == 200 and b'<script' not in html and resp.getheader('Content-Type').startswith('text/html')
    assert api('GET', f'/v1/sessions/{sid}/report?format=xml')[0] == 400


def test_tool_errors_are_reported_as_isError_not_http_errors(api):
    sid = make_session(api)['session_id']
    assert tool(api, 'merge_clusters', session_id=sid, cluster_ids=[0])[0]
    assert tool(api, 'get_cluster_examples', session_id=sid)[0]                  # missing cluster_id
    assert tool(api, 'get_overview', session_id='nope')[0]
    assert tool(api, 'drop_tables', session_id=sid)[0]
    assert tool(api, 'move_responses', session_id=sid, response_numbers=[1], to_cluster=999)[0]
    assert tool(api, 'set_cluster_label', session_id=sid, cluster_id='abc', label='x')[0]


def test_delete_session(api):
    sid = make_session(api)['session_id']
    assert api('DELETE', f'/v1/sessions/{sid}')[1] == {'deleted': True}
    assert api('GET', f'/v1/sessions/{sid}')[0] == 400


def test_body_size_limit(api, monkeypatch):
    monkeypatch.setattr(server, 'MAX_BODY', 50)
    status, out, _ = api('POST', '/v1/sessions', {'texts': ['x' * 200]})
    assert status == 413


def test_internal_errors_do_not_leak(api, monkeypatch):
    def boom(self, body):
        raise RuntimeError('secret internal detail with /home/user/path')
    monkeypatch.setattr(server.App, 'create', boom)
    status, out, _ = api('POST', '/v1/sessions', {'texts': ['a']})
    assert status == 500 and 'secret' not in json.dumps(out)


def test_cli_refuses_network_bind_and_weak_token(monkeypatch, capsys):
    assert server.cli(['--host', '0.0.0.0']) == 2 and 'allow-network' in capsys.readouterr().err
    monkeypatch.setenv('TEXTURR_TOKEN', 'short')
    assert server.cli(['--allow-network', '--host', '0.0.0.0']) == 2 and 'at least 16' in capsys.readouterr().err


# --- file upload -------------------------------------------------------------------------------

def xlsx_bytes(rows):
    import io
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    for r in rows:
        ws.append(r)
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


SHEET = [['Name', 'Comments']] + [['x', f'{k} comment {i}'] for k in DIRS for i in range(6)] + [['x', 'N/A'], ['x', None]]


@pytest.fixture
def srv_port():
    srv = server.make_server(server.App(TOKEN, embed), '127.0.0.1', 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield srv.server_port
    srv.shutdown()


def post_file(port, data, query, token=TOKEN):
    conn = http.client.HTTPConnection('127.0.0.1', port, timeout=10)
    conn.request('POST', '/v1/sessions/from-file?' + query, data, {'Authorization': f'Bearer {token}'})
    r = conn.getresponse(); body = r.read(); conn.close()
    return r.status, (json.loads(body) if body else None)


def get(port, path):
    conn = http.client.HTTPConnection('127.0.0.1', port, timeout=10)
    conn.request('GET', path, headers={'Authorization': f'Bearer {TOKEN}'})
    r = conn.getresponse(); body = r.read(); conn.close()
    return r.status, body, r


def test_upload_xlsx_by_header_name_then_annotated_download(srv_port):
    from openpyxl import load_workbook
    import io
    status, o = post_file(srv_port, xlsx_bytes(SHEET), 'filename=survey.xlsx&column=comments&clusters=4')
    assert status == 201 and o['column'] == 'Comments' and o['groups'] == 4 and o['non_answers'] == 1
    sid = o['session_id']
    status, body, resp = get(srv_port, f'/v1/sessions/{sid}/report?format=annotated')
    assert status == 200 and resp.getheader('Content-Type').startswith('application/vnd.openxmlformats')
    assert 'survey_texturr.xlsx' in resp.getheader('Content-Disposition')
    ws = load_workbook(io.BytesIO(body)).active
    assert ws['C1'].value == 'Comments Cluster' and ws['D1'].value == 'Comments Theme'
    assert ws['B2'].value == 'slow comment 0' and ws['C2'].value is not None and ws['D2'].value   # row alignment kept
    assert ws['D26'].value == 'Non-answer'                         # the 'N/A' row (sheet row 26)
    assert ws['C27'].value is None                                 # the empty cell is untouched
    assert ws['A2'].value == 'x'                                   # original data preserved


def test_upload_csv_and_annotated_csv(srv_port):
    csv_data = ('Name,Comments\n' + '\n'.join(f'x,{k} comment {i}' for k in DIRS for i in range(6))).encode()
    status, o = post_file(srv_port, csv_data, 'filename=s.csv&column=B&clusters=4')
    assert status == 201 and o['groups'] == 4
    status, body, resp = get(srv_port, f"/v1/sessions/{o['session_id']}/report?format=annotated")
    lines = body.decode().splitlines()
    assert status == 200 and lines[0] == 'Name,Comments,Comments Cluster,Comments Theme' and len(lines) == 25


def test_upload_errors(srv_port, monkeypatch):
    good = xlsx_bytes(SHEET)
    assert post_file(srv_port, good, 'filename=s.xlsx')[0] == 400                                   # no column
    status, out = post_file(srv_port, good, 'filename=s.xlsx&column=Nope')
    assert status == 400 and 'Available' in out['error']
    assert post_file(srv_port, b'not a zip', 'filename=s.xlsx&column=A')[0] == 400
    assert post_file(srv_port, b'', 'filename=s.xlsx&column=A')[0] == 400
    assert post_file(srv_port, good, 'filename=s.xlsx&column=A&sheet=Missing')[0] == 400
    assert post_file(srv_port, good, 'filename=s.xlsx&column=Comments&clusters=zero')[0] == 400
    assert post_file(srv_port, good, 'filename=s.xlsx&column=Name&header_row=1')[0] == 201
    assert post_file(srv_port, good, 'filename=s.xlsx&column=Comments', token='bad')[0] == 401
    monkeypatch.setattr(server, 'MAX_UNZIPPED', 100)                                                # zip-bomb guard
    status, out = post_file(srv_port, good, 'filename=s.xlsx&column=Comments')
    assert status == 400 and 'unreasonable' in out['error']


def test_annotated_requires_an_uploaded_file(srv_port):
    conn = http.client.HTTPConnection('127.0.0.1', srv_port, timeout=10)
    conn.request('POST', '/v1/sessions', json.dumps({'texts': TEXTS, 'clusters': 4}), {'Authorization': f'Bearer {TOKEN}'})
    sid = json.loads(conn.getresponse().read())['session_id']
    status, body, _ = get(srv_port, f'/v1/sessions/{sid}/report?format=annotated')
    assert status == 400 and b'uploaded file' in body


def test_hostile_filename_is_neutralized(srv_port):
    status, o = post_file(srv_port, xlsx_bytes(SHEET), 'filename=../../etc/pass"wd%0d%0a.xlsx&column=Comments&clusters=4')
    assert status == 201 and '/' not in o['filename'] and '"' not in o['filename'] and '\n' not in o['filename']
