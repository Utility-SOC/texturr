"""Shareable outputs: a self-contained HTML report and an annotated copy of the input."""

import html
import os

import pandas as pd

CSS = """
:root{--bg:#fff;--fg:#1a1a1a;--muted:#555;--card:#f6f6f4;--line:#d4d4d0;--accent:#0b5cad;--bar:#0b5cad33}
@media (prefers-color-scheme:dark){:root{--bg:#161616;--fg:#ececec;--muted:#b0b0b0;--card:#202020;--line:#3a3a3a;--accent:#7db7f5;--bar:#7db7f544}}
body{font:16px/1.5 system-ui,sans-serif;background:var(--bg);color:var(--fg);margin:0 auto;max-width:60rem;padding:1rem 1rem 3rem}
h1{font-size:1.6rem}h2{margin-top:2.5rem;font-size:1.3rem}h3{margin:0 0 .25rem;font-size:1.1rem}
.note{border:1px solid var(--line);background:var(--card);padding:.6rem .9rem;border-radius:6px;color:var(--muted)}
.card{border:1px solid var(--line);background:var(--card);border-radius:8px;padding:1rem;margin:1rem 0}
.meta{color:var(--muted);font-size:.9rem}
.bar{background:var(--bar);height:.5rem;border-radius:4px;margin:.4rem 0}.bar i{display:block;height:100%;background:var(--accent);border-radius:4px}
blockquote{margin:.4rem 0;padding-left:.8rem;border-left:3px solid var(--line)}
.action{font-weight:600}
"""


def _e(value):
    return html.escape(str(value), quote=True)


def render_html(source, results, generated_by=''):
    """results: list of (column_name, rows, total_responses). Every value is HTML-escaped; no scripts or external assets."""
    parts = [f"<!doctype html><html lang=en><head><meta charset=utf-8>"
             f"<meta name=viewport content='width=device-width,initial-scale=1'>"
             f"<meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; style-src 'unsafe-inline'\">"
             f"<title>texturr report: {_e(os.path.basename(source))}</title><style>{CSS}</style></head><body>"
             f"<main><h1>texturr report</h1><p class=meta>Source: {_e(os.path.basename(source))}"
             f"{' &middot; ' + _e(generated_by) if generated_by else ''}</p>"
             "<p class=note>This report contains verbatim responses from the source data. "
             "Handle it at the same classification as the source.</p>"]
    for column, rows, total in results:
        parts.append(f"<section><h2>{_e(column)} <span class=meta>({total} responses, {len(rows)} groups)</span></h2>")
        for r in sorted(rows, key=lambda r: -r['Size']):
            pct = 100 * r['Size'] / max(1, total)
            title = r['Label'] or f"Group {r['Cluster']}"
            parts.append(
                f"<article class=card><h3>{_e(title)}</h3>"
                f"<p class=meta>{r['Size']} responses ({pct:.0f}%)</p>"
                f"<div class=bar role=img aria-label='{pct:.0f} percent of responses'><i style='width:{pct:.1f}%'></i></div>")
            if r['Summary']:
                parts.append(f"<p>{_e(r['Summary'])}</p>")
            if r['Suggested Action']:
                parts.append(f"<p class=action>Suggested action: <span style='font-weight:400'>{_e(r['Suggested Action'])}</span></p>")
            if r['Keyphrases']:
                parts.append(f"<p class=meta>Keyphrases: {_e(r['Keyphrases'])}</p>")
            for q in str(r['Representative Responses']).split(' | '):
                if q:
                    parts.append(f"<blockquote>{_e(q)}</blockquote>")
            parts.append("</article>")
        parts.append("</section>")
    parts.append("</main></body></html>")
    return ''.join(parts)


def write_html(path, source, results, generated_by=''):
    with open(path, 'w', encoding='utf-8') as f:
        f.write(render_html(source, results, generated_by))


def cell_text(value):
    """Plain text for a spreadsheet cell."""
    return str(value)


def annotate_workbook(source, sheet, annotations, output, header_row=1):
    """Copy an .xlsx and add `<header> Cluster` / `<header> Theme` columns beside the data.

    annotations: list of dicts with 'header', 'positions' (0-based sheet row per response),
    'clusters', and 'labels' (cluster id -> theme text). Strings are stored as text so a
    label beginning with '=' can never become a formula. New headers go on `header_row`
    (1-based); with no header row (0) they are omitted rather than overwrite data.
    """
    from openpyxl import load_workbook
    wb = load_workbook(source)
    ws = wb[sheet] if sheet else wb.active
    col = ws.max_column + 1
    for a in annotations:
        if header_row:
            ws.cell(row=header_row, column=col, value=f"{a['header']} Cluster")
            ws.cell(row=header_row, column=col + 1, value=f"{a['header']} Theme")
        for pos, cid in zip(a['positions'], a['clusters']):
            c1 = ws.cell(row=pos + 1, column=col, value=int(cid))
            c2 = ws.cell(row=pos + 1, column=col + 1)
            c2.value = cell_text(a['labels'].get(int(cid), ''))
            c2.data_type = 's'
        col += 2
    wb.save(output)


def annotate_csv(source, annotations, output, csv_safe, header_row=1):
    """CSV/TSV equivalent: appends cluster and theme columns to the input rows."""
    sep = '\t' if source.lower().endswith(('.tsv', '.tab')) else ','
    df = pd.read_csv(source, header=None, sep=sep, dtype=str, keep_default_na=False, skip_blank_lines=False)
    for a in annotations:
        cl = [''] * len(df)
        th = [''] * len(df)
        if header_row:
            cl[header_row - 1] = f"{a['header']} Cluster"
            th[header_row - 1] = f"{a['header']} Theme"
        for pos, cid in zip(a['positions'], a['clusters']):
            cl[pos] = str(int(cid))
            th[pos] = csv_safe(a['labels'].get(int(cid), ''))
        df[f"{a['header']} Cluster"] = cl
        df[f"{a['header']} Theme"] = th
    df.to_csv(output, index=False, header=False, sep=sep)
